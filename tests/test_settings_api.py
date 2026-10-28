import io
import json

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import create_app
from app.provider import ModelConfig, ModelProvider


@pytest.fixture
def active(tmp_path, monkeypatch):
    for name in ('RAG_BASE_URL', 'RAG_MODEL', 'RAG_API_KEY', 'ZHIXU_SEED_DEMO'):
        monkeypatch.delenv(name, raising=False)
    with TestClient(create_app(db_path=tmp_path / 'data' / 'knowledge.db', seed_demo=False)) as client:
        yield client


def test_settings_can_be_saved_without_restart_and_never_return_key(active):
    body = {'base_url': 'https://model.example/v1', 'model': 'demo', 'api_key': 'private-test-value'}
    saved = active.put('/api/settings/model', json=body)
    assert saved.status_code == 200, saved.text
    assert saved.json()['enabled'] is True
    assert saved.json()['has_api_key'] is True
    assert 'private-test-value' not in saved.text
    assert 'private-test-value' not in active.get('/api/settings/model').text
    assert active.get('/api/health').json()['model']['name'] == 'demo'
    assert active.put('/api/settings/model', json={**body, 'api_key': None, 'model': 'next'}).json()['has_api_key']
    assert not active.put('/api/settings/model', json={**body, 'clear_api_key': True}).json()['has_api_key']


def test_api_validation_does_not_echo_secret_values(active):
    secret = 'never-echo-this'
    response = active.put('/api/settings/model', json={'base_url': 'http://external.example/v1', 'model': 'm', 'api_key': secret})
    assert response.status_code == 400
    assert secret not in response.text
    response = active.put('/api/settings/model', json={'base_url': 'https://example.com/v1', 'model': 'm', 'api_key': {'secret': secret}})
    assert response.status_code == 422
    assert secret not in response.text


def test_settings_reject_cross_origin_change(active):
    assert active.put('/api/settings/model', json={'base_url': 'http://localhost:1234/v1', 'model': 'm'}, headers={'origin': 'https://evil.example'}).status_code == 403


@pytest.mark.parametrize('host', ['[::1]:18765', '127.0.0.2:18765'])
def test_all_cli_loopback_hosts_are_accepted(tmp_path, monkeypatch, host):
    monkeypatch.setenv('ZHIXU_HOST', '::1' if host.startswith('[') else '127.0.0.2')
    with TestClient(create_app(db_path=tmp_path / 'test.db', seed_demo=False)) as client:
        assert client.get('/api/health', headers={'host': host}).status_code == 200


def test_injected_model_is_locked_and_probe_is_synthetic(tmp_path):
    requests = []
    def handler(request):
        requests.append(json.loads(request.content))
        return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps({'answer': '连接测试代号为星舟。[1]', 'source_ids': [1], 'status': 'answered'}, ensure_ascii=False)}}]})
    model = ModelProvider(ModelConfig('http://127.0.0.1:1234/v1', 'local', ''), transport=httpx.MockTransport(handler))
    with TestClient(create_app(db_path=tmp_path / 'test.db', provider=model, seed_demo=False)) as client:
        client.post('/api/documents', files={'file': ('private.txt', '私人文件内容不发给连接测试'.encode())})
        assert client.get('/api/settings/model').json()['editable'] is False
        assert client.put('/api/settings/model', json={'base_url': 'http://localhost:1234/v1', 'model': 'next'}).status_code == 409
        response = client.post('/api/settings/model/test')
        assert response.status_code == 200, response.text
        assert response.json()['ok'] is True
        assert len(requests) == 1
        assert '星舟' in requests[0]['messages'][1]['content']
        assert '私人文件' not in str(requests)


def test_failed_generation_returns_honest_extractive_fallback(tmp_path):
    model = ModelProvider(ModelConfig('http://localhost:1234/v1', 'local', ''), transport=httpx.MockTransport(lambda request: httpx.Response(500)))
    with TestClient(create_app(db_path=tmp_path / 'test.db', provider=model, seed_demo=False)) as client:
        client.post('/api/documents', files={'file': ('规则.txt', '平台试用期为 14 天。'.encode())})
        response = client.post('/api/ask', json={'question': '平台试用期几天？', 'mode': 'llm'})
        assert response.status_code == 200
        body = response.json()
        assert body['mode'] == 'extractive'
        assert '14' in body['answer']
        assert '500' in body['generation_error']
        assert '摘录' in body['warning']
        assert client.post('/api/settings/model/test').status_code == 502


def test_docx_upload_can_be_queried_and_locations_are_honest(active):
    document = Document()
    document.add_paragraph('平台试用期为 14 天。')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '项目数量'
    table.cell(0, 1).text = '3 个'
    stream = io.BytesIO()
    document.save(stream)
    response = active.post('/api/documents', files={'file': ('规则.docx', stream.getvalue())})
    assert response.status_code == 200, response.text
    original = active.get('/api/documents/' + response.json()['id']).json()
    assert original['kind'] == 'docx'
    assert any('表 1' in chunk['location'] for chunk in original['chunks'])
    assert all('页' not in chunk['location'] for chunk in original['chunks'])
    answer = active.post('/api/ask', json={'question': '平台试用期几天？'}).json()
    assert '14' in answer['answer']


def test_default_workspace_is_empty_and_demo_resources_are_available(tmp_path, monkeypatch):
    monkeypatch.setenv('ZHIXU_DATA_DIR', str(tmp_path / 'workspace'))
    monkeypatch.delenv('ZHIXU_SEED_DEMO', raising=False)
    with TestClient(create_app()) as client:
        assert client.get('/api/documents').json()['documents'] == []
        assert client.get('/api/health').json()['version'] == '1.2.0'
        assert client.get('/api/config-template').status_code == 200
        assert len(client.post('/api/demo/restore').json()['documents']) == 8
    assert (tmp_path / 'workspace' / 'knowledge.db').exists()


def test_api_docs_use_packaged_assets_and_no_inline_bootstrap(active):
    response = active.get('/api/docs')
    assert response.status_code == 200
    assert 'https://' not in response.text
    assert '<script>' not in response.text
    assert '/static/openapi.js' in response.text
    assert active.get('/static/swagger/swagger-ui-bundle.js').status_code == 200
    assert active.get('/openapi.json').json()['info']['version'] == '1.2.0'
