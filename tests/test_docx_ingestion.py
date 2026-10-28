import base64
import io
import random
import struct
import zipfile

from docx import Document
import pytest

from app.docx_ingestion import extract_docx_sections


def _save(document):
    target = io.BytesIO()
    document.save(target)
    return target.getvalue()


def _document(text='正文说明。'):
    document = Document()
    document.add_paragraph(text)
    return _save(document)


def _rewrite(content, replacements=None, extras=(), remove=(), compression=zipfile.ZIP_DEFLATED):
    target = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(target, 'w', compression=compression) as output:
        for entry in original.infolist():
            if entry.filename not in remove:
                output.writestr(entry.filename, (replacements or {}).get(entry.filename, original.read(entry)))
        for name, payload in extras:
            output.writestr(name, payload)
    return target.getvalue()


def test_paragraphs_titles_and_table_rows_preserve_body_order():
    document = Document()
    document.add_heading('中文资料标题', level=0)
    document.add_paragraph('开始的正文。')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '项目'
    table.cell(0, 1).text = '说明'
    table.cell(1, 0).text = '申请'
    table.cell(1, 1).text = '先向主管提交。'
    document.add_paragraph('')
    document.add_paragraph('表格之后的正文。')
    assert extract_docx_sections(_save(document)) == [
        ('第 1 段', '中文资料标题'), ('第 2 段', '开始的正文。'),
        ('表 1 · 第 1 行', '项目\t说明'), ('表 1 · 第 2 行', '申请\t先向主管提交。'),
        ('第 4 段', '表格之后的正文。'),
    ]


def test_multiple_tables_have_distinct_locations_and_empty_rows_are_skipped():
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(1, 0).text = '第一张表。'
    document.add_paragraph('中间说明。')
    document.add_table(rows=1, cols=1).cell(0, 0).text = '第二张表。'
    assert extract_docx_sections(_save(document)) == [
        ('表 1 · 第 2 行', '第一张表。'), ('第 1 段', '中间说明。'), ('表 2 · 第 1 行', '第二张表。'),
    ]


def test_only_body_is_extracted_not_headers_and_footers():
    document = Document()
    document.sections[0].header.paragraphs[0].text = '页眉内容。'
    document.sections[0].footer.paragraphs[0].text = '页脚内容。'
    document.add_paragraph('正文内容。')
    assert extract_docx_sections(_save(document)) == [('第 1 段', '正文内容。')]


def test_empty_document_is_clear_error():
    with pytest.raises(ValueError, match='文字|正文'):
        extract_docx_sections(_save(Document()))


def test_image_only_document_is_not_reported_as_success():
    document = Document()
    image = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j00kAAAAASUVORK5CYII=')
    document.add_picture(io.BytesIO(image))
    with pytest.raises(ValueError, match='图片|文字|OCR'):
        extract_docx_sections(_save(document))


@pytest.mark.parametrize('content', [b'', b'not a zip', b'PK\x03\x04broken', b'\xd0\xcf\x11\xe0legacy-word'])
def test_corrupt_or_legacy_content_is_value_error(content):
    with pytest.raises(ValueError):
        extract_docx_sections(content)


def test_compressed_upload_size_is_limited():
    with pytest.raises(ValueError, match='10'):
        extract_docx_sections(b'x' * (10 * 1024 * 1024 + 1))


def test_missing_main_document_is_rejected():
    with pytest.raises(ValueError, match='document.xml|正文'):
        extract_docx_sections(_rewrite(_document(), remove={'word/document.xml'}))


def test_single_zip_member_has_eight_mib_limit():
    payload = b'a' * (8 * 1024 * 1024 + 1)
    with pytest.raises(ValueError, match='8 MiB|单个'):
        extract_docx_sections(_rewrite(_document(), extras=[('custom/oversized.bin', payload)]))


def test_zip_members_share_total_decompression_budget():
    payload = bytes(range(256)) * (5 * 1024 * 1024 // 256)
    content = _rewrite(_document(), extras=[(f'custom/part-{index}.bin', payload) for index in range(5)])
    assert len(content) < 1024 * 1024
    with pytest.raises(ValueError, match='累计|总解压'):
        extract_docx_sections(content)


def test_archive_entry_count_is_limited():
    content = _rewrite(_document(), extras=[(f'custom/part-{index}.bin', b'x') for index in range(2001)])
    with pytest.raises(ValueError, match='2000|数量'):
        extract_docx_sections(content)


@pytest.mark.parametrize('name', ['../escape', 'word/../escape', '..\\escape', '/absolute', 'C:/escape', 'word/./document.xml'])
def test_unsafe_member_paths_are_rejected(name):
    with pytest.raises(ValueError, match='路径'):
        extract_docx_sections(_rewrite(_document(), extras=[(name, b'unsafe')]))


def test_duplicate_member_names_are_rejected():
    with pytest.warns(UserWarning, match='Duplicate name'):
        content = _rewrite(_document(), extras=[('word/document.xml', b'duplicate')])
    with pytest.raises(ValueError, match='重复'):
        extract_docx_sections(content)


def test_encrypted_zip_flag_is_rejected_before_parsing():
    content = bytearray(_document())
    local = content.index(b'PK\x03\x04')
    central = content.index(b'PK\x01\x02')
    struct.pack_into('<H', content, local + 6, struct.unpack_from('<H', content, local + 6)[0] | 1)
    struct.pack_into('<H', content, central + 8, struct.unpack_from('<H', content, central + 8)[0] | 1)
    with pytest.raises(ValueError, match='加密'):
        extract_docx_sections(bytes(content))


def test_extreme_compression_ratio_is_rejected():
    content = _rewrite(_document(), extras=[('custom/extreme.bin', b'\0' * (2 * 1024 * 1024))])
    with pytest.raises(ValueError, match='压缩比'):
        extract_docx_sections(content)


def test_normal_compressed_repeated_chinese_text_is_accepted():
    text = '这是一段正常的中文资料，包含重复出现的专业说明。' * 4000
    assert extract_docx_sections(_document(text)) == [('第 1 段', text)]


def test_macro_payload_is_rejected_even_if_main_content_is_docx():
    with pytest.raises(ValueError, match='宏'):
        extract_docx_sections(_rewrite(_document(), extras=[('word/vbaProject.bin', b'macro')]))


@pytest.mark.parametrize('encoding', ['utf-8', 'utf-16'])
def test_dtd_and_external_entities_are_rejected_before_library_parse(encoding, monkeypatch):
    import app.docx_ingestion as module
    content = _document()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        original = archive.read('word/document.xml').decode('utf-8')
    body = original.split('?>', 1)[1].replace('正文说明。', '&external;')
    xml = (f'<?xml version="1.0" encoding="{encoding}"?>'
           '<!DOCTYPE w:document [<!ENTITY external SYSTEM "https://example.invalid/private">]>' + body).encode(encoding)
    calls = []
    monkeypatch.setattr(module, 'Document', lambda *args, **kwargs: calls.append(args))
    with pytest.raises(ValueError, match='DTD|实体'):
        extract_docx_sections(_rewrite(content, replacements={'word/document.xml': xml}))
    assert calls == []


def test_external_relationship_does_not_require_fetching_network_content():
    document = Document()
    document.add_paragraph('本机正文。')
    document.part.relate_to('https://example.invalid/never-fetch',
                           'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                           is_external=True)
    assert extract_docx_sections(_save(document)) == [('第 1 段', '本机正文。')]


def test_nested_table_content_remains_in_outer_row_in_cell_order():
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = '单元格开始。'
    cell.add_table(rows=1, cols=1).cell(0, 0).text = '内嵌表格说明。'
    cell.add_paragraph('单元格结束。')
    assert extract_docx_sections(_save(document)) == [
        ('表 1 · 第 1 行', '单元格开始。\n内嵌表格说明。\n单元格结束。'),
    ]


def test_corrupt_unreferenced_zip_member_is_still_rejected():
    content = bytearray(_rewrite(_document(), extras=[('custom/check.bin', b'crc-sentinel-first')],
                                 compression=zipfile.ZIP_STORED))
    content[content.index(b'crc-sentinel-first')] ^= 1
    with pytest.raises(ValueError, match='损坏'):
        extract_docx_sections(bytes(content))


def test_non_office_zip_compression_is_rejected_without_decompressing():
    content = _rewrite(_document(), compression=zipfile.ZIP_BZIP2)
    with pytest.raises(ValueError, match='压缩方式'):
        extract_docx_sections(content)


def test_large_grid_span_does_not_duplicate_a_physical_cell():
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    text = '甲' * 1000
    cell.text = text
    cell._tc.get_or_add_tcPr().get_or_add_gridSpan().val = 5000
    content = _save(document)
    assert len(content) < 50 * 1024
    sections = extract_docx_sections(content)
    assert len(sections) == 1
    assert sections[0][0] == '表 1 · 第 1 行'
    assert len(sections[0][1]) == len(text)
    assert sections[0][1] == text


def test_nested_grid_spans_do_not_multiply_inner_content():
    document = Document()
    outer = document.add_table(rows=1, cols=1).cell(0, 0)
    inner = outer.add_table(rows=1, cols=1).cell(0, 0)
    text = '内嵌说明' * 100
    inner.text = text
    for cell in (outer, inner):
        cell._tc.get_or_add_tcPr().get_or_add_gridSpan().val = 30
    sections = extract_docx_sections(_save(document))
    assert len(sections[0][1]) == len(text)
    assert sections[0][1] == text


def test_horizontal_and_vertical_merges_only_emit_original_cell_content():
    document = Document()
    horizontal = document.add_table(rows=1, cols=2)
    horizontal.cell(0, 0).merge(horizontal.cell(0, 1)).text = '横向合并。'
    vertical = document.add_table(rows=2, cols=1)
    vertical.cell(0, 0).merge(vertical.cell(1, 0)).text = '纵向合并。'
    assert extract_docx_sections(_save(document)) == [
        ('表 1 · 第 1 行', '横向合并。'), ('表 2 · 第 1 行', '纵向合并。'),
    ]


def test_extracted_text_has_a_cumulative_budget_before_sections_return():
    generator = random.Random(0)
    block = ''.join(generator.choices('abcdefghijklmnopqrstuvwxyz0123456789', k=1024))
    text = (block * 1000)[:1_000_001]
    document = Document()
    document.add_paragraph(text)
    document.add_paragraph(text)
    with pytest.raises(ValueError, match='文字|字符'):
        extract_docx_sections(_save(document))


def test_physical_body_structure_has_a_bounded_node_count(monkeypatch):
    import app.docx_ingestion as module
    monkeypatch.setattr(module, 'MAX_DOCX_BODY_ELEMENTS', 50, raising=False)
    document = Document()
    for _ in range(30):
        document.add_paragraph('结构预算。')
    with pytest.raises(ValueError, match='结构|节点'):
        extract_docx_sections(_save(document))


def test_nested_table_depth_is_limited():
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    for _ in range(16):
        cell = cell.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = '最深处文字。'
    with pytest.raises(ValueError, match='嵌套|深度'):
        extract_docx_sections(_save(document))
